"""Vakter för enhetsjournalen (assets/downloads/two-minute-debrief-journal.xlsx).

Journalen byggs av exports/build-two-minute-debrief-journal.py från
journal_example i two-minute-debrief/debrief-en.json. Den öppnas i Excel och
kan importeras i Google Sheets. Inga formler eller makron.
"""
import importlib.util
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "assets/downloads/two-minute-debrief-journal.xlsx"
DATA = ROOT / "two-minute-debrief/debrief-en.json"
BUILD = ROOT / "exports/build-two-minute-debrief-journal.py"
HEADERS = ["Date", "Group", "Goal", "Evidence", "Surprise", "Next", "Change I made", "Check result"]


def rows():
    return json.loads(DATA.read_text(encoding="utf-8"))["journal_example"]


def test_sheets():
    assert load_workbook(XLSX).sheetnames == ["Journal", "Example", "How to use"]


def test_journal_sheet_is_empty_with_header_and_frozen_row():
    ws = load_workbook(XLSX)["Journal"]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.freeze_panes == "A2"
    assert all(c.value is None for c in ws[2])


def test_example_sheet_has_journal_example_rows():
    ws = load_workbook(XLSX)["Example"]
    assert [c.value for c in ws[1]] == HEADERS
    first, last = rows()[0], rows()[-1]
    assert [ws.cell(2, i).value for i in range(1, 9)] == [
        first["date"], first["group"], first["goal"], first["evidence"], first["surprise"],
        first["next"], first["change"], first["check"]]
    assert ws.cell(1 + len(rows()), 8).value == last["check"]


def test_no_formulas():
    for ws in load_workbook(XLSX):
        for row in ws.iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith("=")), c.coordinate


def test_how_to_use_mentions_page_rollup_and_privacy():
    ws = load_workbook(XLSX)["How to use"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "choosewise.education/two-minute-debrief" in text
    assert "roll-up" in text
    assert "student names" in text


def test_build_function_uses_the_json():
    spec = importlib.util.spec_from_file_location("journal", BUILD)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.build_workbook(rows())["Example"].cell(2, 3).value == rows()[0]["goal"]
