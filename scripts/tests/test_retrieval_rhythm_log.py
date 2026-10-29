"""Vakter för loggmallen (assets/downloads/retrieval-rhythm-log.xlsx).

Mallen byggs av exports/build-retrieval-rhythm-log.py från exempelloggen i
retrieval-rhythm/rhythm-en.json. Den öppnas i Excel och kan importeras i
Google Sheets. Inga formler eller makron.
"""
import importlib.util
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "assets/downloads/retrieval-rhythm-log.xlsx"
DATA = ROOT / "retrieval-rhythm/rhythm-en.json"
BUILD = ROOT / "exports/build-retrieval-rhythm-log.py"
HEADERS = ["Date", "Topic", "Key points", "Notes"]


def rows():
    return json.loads(DATA.read_text(encoding="utf-8"))["example_log"]


def test_workbook_has_both_sheets():
    wb = load_workbook(XLSX)
    assert wb.sheetnames == ["Log", "How to use"]


def test_log_header_and_example_rows():
    ws = load_workbook(XLSX)["Log"]
    assert [c.value for c in ws[1]] == HEADERS
    assert ws.freeze_panes == "A2"
    first = rows()[0]
    assert [ws.cell(2, i).value for i in (1, 2, 3)] == [first["date"], first["topic"], first["key_points"]]
    assert ws.cell(1 + len(rows()), 2).value == rows()[-1]["topic"]


def test_no_formulas():
    wb = load_workbook(XLSX)
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith("=")), c.coordinate


def test_how_to_use_mentions_the_prompt_and_privacy():
    ws = load_workbook(XLSX)["How to use"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "choosewise.education/retrieval-rhythm" in text
    assert "student names" in text


def test_build_function_uses_the_json():
    spec = importlib.util.spec_from_file_location("log", BUILD)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    wb = mod.build_workbook(rows())
    assert wb["Log"].cell(2, 2).value == rows()[0]["topic"]
