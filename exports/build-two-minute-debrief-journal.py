#!/usr/bin/env python3
"""Build the Two-Minute Debrief unit journal (.xlsx) from two-minute-debrief/debrief-en.json.

Sheet "Journal": header Date | Group | Goal | Evidence | Surprise | Next |
Change I made | Check result, and empty rows to fill in. Sheet "Example": the
same with the example unit. Sheet "How to use": a few short lines.
Plain values only: no formulas, no macros. Opens in Excel; Google Sheets
imports it via File → Import.

Usage:
    /usr/bin/python3 exports/build-two-minute-debrief-journal.py
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "two-minute-debrief/debrief-en.json"
DST = ROOT / "assets/downloads/two-minute-debrief-journal.xlsx"

HEADERS = ["Date", "Group", "Goal", "Evidence", "Surprise", "Next", "Change I made", "Check result"]
KEYS = ["date", "group", "goal", "evidence", "surprise", "next", "change", "check"]
WIDTHS = {"A": 12, "B": 22, "C": 32, "D": 40, "E": 30, "F": 26, "G": 32, "H": 28}
EMPTY_ROWS = 60
# Sajtens blå (tokens.css --color-accent) som rubrikrad, vit text.
HEADER_FILL = PatternFill("solid", fgColor="0B3A6F")
HOW_TO_USE = [
    "Two-Minute Debrief — unit journal",
    "1. After a lesson that didn't go as planned, answer the four questions at choosewise.education/two-minute-debrief and save or copy the prompt.",
    "2. Each saved debrief has a “Copy journal row” button. Paste the row on the Journal sheet.",
    "3. After the next lesson, fill in Change I made and Check result. The check result is also the evidence in your next debrief.",
    "4. Once a week, paste the rows into the weekly roll-up on the same page to see patterns across lessons.",
    "Only the group goes in this journal. Never add student names or individual results.",
]


def _sheet(ws, rows: list) -> None:
    """Header row, the given rows, then empty rows to fill in."""
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF")
        cell.fill = HEADER_FILL
    for r in rows:
        ws.append([r[k] for k in KEYS])
    for _ in range(EMPTY_ROWS):
        ws.append([None] * len(HEADERS))
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def build_workbook(rows: list) -> Workbook:
    # Journalbladet är tomt för läraren; exemplet ligger på ett eget blad.
    wb = Workbook()
    journal = wb.active
    journal.title = "Journal"
    _sheet(journal, [])
    _sheet(wb.create_sheet("Example"), rows)

    how = wb.create_sheet("How to use")
    for line in HOW_TO_USE:
        how.append([line])
    how.column_dimensions["A"].width = 110
    return wb


def main() -> None:
    rows = json.loads(DATA.read_text(encoding="utf-8"))["journal_example"]
    DST.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(rows).save(DST)
    print(f"Built {DST.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
