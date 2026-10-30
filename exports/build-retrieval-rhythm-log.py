#!/usr/bin/env python3
"""Build the Retrieval Rhythm topic log template (.xlsx) from retrieval-rhythm/rhythm-en.json.

Sheet "Log": header Date | Topic | Key points | Notes (not for the prompt), and
empty rows to fill in. Sheet "Example": the same with the example rows.
Sheet "How to use": a few short lines.
Plain values only: no formulas, no macros. Opens in Excel; Google Sheets
imports it via File → Import.

Usage:
    /usr/bin/python3 exports/build-retrieval-rhythm-log.py
"""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "retrieval-rhythm/rhythm-en.json"
DST = ROOT / "assets/downloads/retrieval-rhythm-log.xlsx"

HEADERS = ["Date", "Topic", "Key points", "Notes (not for the prompt)"]
WIDTHS = {"A": 13, "B": 40, "C": 70, "D": 30}
EMPTY_ROWS = 60
# Sajtens blå (tokens.css --color-accent) som rubrikrad, vit text.
HEADER_FILL = PatternFill("solid", fgColor="0B3A6F")
HOW_TO_USE = [
    "Retrieval Rhythm — topic log",
    "1. After each lesson, add a row on the Log sheet: the date, the topic and two or three key points. The Example sheet shows what it can look like.",
    "2. On Friday, copy columns A to C and paste them into the prompt builder at choosewise.education/retrieval-rhythm. Leave the Notes column out.",
    "3. Copy the prompt into your AI tool, check the answer key, and put the starters on your slides.",
    "Only topics go in this log. Never add student names or student results.",
]


def _sheet(ws, rows: list) -> None:
    """Header row, the given rows, then empty rows to fill in."""
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF")
        cell.fill = HEADER_FILL
    for r in rows:
        ws.append([r["date"], r["topic"], r["key_points"], None])
    for _ in range(EMPTY_ROWS):
        ws.append([None] * len(HEADERS))
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def build_workbook(rows: list) -> Workbook:
    # Loggbladet är tomt för läraren; exemplet ligger på ett eget blad så att
    # det inte följer med in i en prompt i ett annat ämne.
    wb = Workbook()
    log = wb.active
    log.title = "Log"
    _sheet(log, [])
    _sheet(wb.create_sheet("Example"), rows)

    how = wb.create_sheet("How to use")
    for line in HOW_TO_USE:
        how.append([line])
    how.column_dimensions["A"].width = 110
    return wb


def main() -> None:
    rows = json.loads(DATA.read_text(encoding="utf-8"))["example_log"]
    DST.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(rows).save(DST)
    print(f"Built {DST.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
